"""Offline tests for the optional native Word argument table exporter.

Source records are synthetic fixtures; these tests validate structure and delivery,
not the truth of any argument or the reachability of a source.
"""
from __future__ import annotations

from contextlib import redirect_stderr, redirect_stdout
from copy import deepcopy
import importlib.util
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
from zipfile import ZipFile
from unittest.mock import patch

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Cm


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SCRIPT_DIR = PROJECT_ROOT / "skills" / "一辩稿DOCX输出" / "scripts"
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))
spec = importlib.util.spec_from_file_location("export_argument_table", SCRIPT_DIR / "export_argument_table.py")
if spec is None or spec.loader is None:
    raise RuntimeError("Cannot load argument table exporter")
exporter = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = exporter
spec.loader.exec_module(exporter)

TOPIC = "AI 2.0：是否会改变学习方式？"
SOURCE_URL = "https://example.org/original-record?topic=learning&year=2026#results"
ARGUMENT = {
    "id": "A1",
    "verified": True,
    "categories": ["进攻", "证明"],
    "claim": "离线测试论据核心",
    "material": "结构测试使用的核验事实与案例简述。",
    "usage": "说明材料与本方判断之间的支持关系。",
    "limitation": "仅在来源界定的对象与条件内使用。",
    "verification_note": "核对了公开记录的对象与结果；此项测试不实际访问来源。",
    "sources": [{
        "author": "测试资料机构",
        "title": "离线测试原始记录",
        "date": "未标注日期",
        "url": SOURCE_URL,
    }],
}


class ArgumentTableOutputTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.products = self.root / "custom products"
        self.input_file = self.root / "arguments.json"

    def payload(self):
        second = deepcopy(ARGUMENT)
        second.update(id="B2", categories=["防守"], claim="第二条测试论据")
        return {"arguments": [deepcopy(ARGUMENT), second]}

    def run_export(self, payload=None, *extra, custom_directory=True):
        if payload is not None:
            self.input_file.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        args = [
            "export_argument_table.py", "--topic", TOPIC, "--side", "反方",
            "--input", str(self.input_file),
        ]
        if custom_directory:
            args += ["--products-dir", str(self.products)]
        args += extra
        stdout, stderr = io.StringIO(), io.StringIO()
        with patch.object(sys, "argv", args), redirect_stdout(stdout), redirect_stderr(stderr):
            code = exporter.main()
        return code, stdout.getvalue(), stderr.getvalue()

    def output_path(self, products=None):
        return (products or self.products) / f"{TOPIC}-反方-论据资料表.docx"

    def assert_no_output(self):
        self.assertFalse(self.products.exists())
        self.assertEqual(list(self.root.rglob("*.docx")), [])

    def test_groups_categories_in_order_and_preserves_same_id_and_sources(self):
        code, stdout, stderr = self.run_export(self.payload())
        self.assertEqual(code, 0, stderr)
        output = self.output_path()
        self.assertIn(str(output.resolve()), stdout)
        self.assertIn("论据资料表已保存", stdout)
        self.assertIn(str(self.products.resolve()), stdout)
        self.assertNotIn("是否生成", stdout)
        doc = Document(output)
        headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
        self.assertEqual(headings, ["进攻", "防守", "证明"])
        self.assertEqual(len(doc.tables), 3)
        self.assertEqual([table.rows[1].cells[0].text for table in doc.tables], ["A1", "B2", "A1"])
        for table in doc.tables:
            self.assertEqual([cell.text for cell in table.rows[0].cells], list(exporter.HEADERS))
            row = table.rows[1]
            self.assertIn(SOURCE_URL, row.cells[4].text)
            self.assertIn("测试资料机构", row.cells[4].text)
            self.assertIn("离线测试原始记录", row.cells[4].text)
            self.assertIn("未标注日期", row.cells[4].text)
            self.assertIn(ARGUMENT["usage"], row.cells[2].text)
            self.assertIn(ARGUMENT["limitation"], row.cells[3].text)
            self.assertIn(ARGUMENT["verification_note"], row.cells[4].text)
        self.assertIn(TOPIC, doc.core_properties.title)

    def test_full_source_urls_are_external_word_hyperlinks_and_keep_wrapping(self):
        self.assertEqual(self.run_export(self.payload())[0], 0)
        namespaces = {
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
            "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
            "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
        }
        with ZipFile(self.output_path()) as archive:
            document_xml = ET.fromstring(archive.read("word/document.xml"))
            relationships_xml = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
        relationships = {
            item.attrib["Id"]: item.attrib
            for item in relationships_xml.findall("rel:Relationship", namespaces)
        }
        link_paragraphs = [
            paragraph for paragraph in document_xml.findall(".//w:p", namespaces)
            if paragraph.find("w:hyperlink", namespaces) is not None
        ]
        self.assertEqual(len(link_paragraphs), 3)
        for paragraph in link_paragraphs:
            hyperlink = paragraph.find("w:hyperlink", namespaces)
            relationship = relationships[hyperlink.attrib[qn("r:id")]]
            self.assertEqual(relationship["Target"], SOURCE_URL)
            self.assertEqual(relationship["TargetMode"], "External")
            self.assertTrue(relationship["Type"].endswith("/hyperlink"))
            self.assertEqual("".join(node.text or "" for node in hyperlink.findall(".//w:t", namespaces)), SOURCE_URL)
            self.assertEqual(paragraph.find("w:pPr/w:wordWrap", namespaces).attrib[qn("w:val")], "0")

    def test_uses_a4_landscape_five_native_columns_and_repeating_header(self):
        self.assertEqual(self.run_export(self.payload())[0], 0)
        doc = Document(self.output_path())
        section = doc.sections[0]
        self.assertEqual(section.orientation, WD_ORIENT.LANDSCAPE)
        self.assertAlmostEqual(section.page_width / Cm(1), 29.7, places=2)
        self.assertAlmostEqual(section.page_height / Cm(1), 21, places=2)
        for table in doc.tables:
            self.assertEqual(len(table.columns), 5)
            self.assertFalse(table.autofit)
            self.assertIsNotNone(table.rows[0]._tr.trPr.find(qn("w:tblHeader")))
            self.assertTrue(all(cell.width > 0 for cell in table.rows[0].cells))
            for cell in table.rows[0].cells:
                self.assertTrue(cell.paragraphs[0].runs[0].bold)

    def test_missing_categories_are_explicitly_empty_without_invented_rows(self):
        first = deepcopy(ARGUMENT)
        first["categories"] = ["证明"]
        code, _, stderr = self.run_export({"arguments": [first]})
        self.assertEqual(code, 0, stderr)
        doc = Document(self.output_path())
        self.assertEqual([p.text for p in doc.paragraphs].count(exporter.NO_MATERIAL_MESSAGE), 2)
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(len(doc.tables[0].rows), 2)
        self.assertEqual(doc.tables[0].rows[1].cells[0].text, "A1")

    def test_rejects_unverified_or_non_boolean_verification_flags(self):
        for flag in (False, "true", 1, None):
            with self.subTest(flag=flag):
                payload = {"arguments": [deepcopy(ARGUMENT)]}
                payload["arguments"][0]["verified"] = flag
                code, stdout, stderr = self.run_export(payload)
                self.assertEqual(code, 2, stderr)
                self.assertNotIn("已保存", stdout)
                self.assert_no_output()

    def test_rejects_each_missing_required_field_or_empty_text(self):
        for field in ("id", "verified", "categories", *exporter.TEXT_FIELDS, "sources"):
            with self.subTest(missing=field):
                payload = {"arguments": [deepcopy(ARGUMENT)]}
                del payload["arguments"][0][field]
                self.assertEqual(self.run_export(payload)[0], 2)
                self.assert_no_output()
        for field in exporter.TEXT_FIELDS:
            with self.subTest(blank=field):
                payload = {"arguments": [deepcopy(ARGUMENT)]}
                payload["arguments"][0][field] = " \n\t"
                self.assertEqual(self.run_export(payload)[0], 2)
                self.assert_no_output()
        for field in exporter.SOURCE_FIELDS:
            with self.subTest(missing_source_field=field):
                payload = {"arguments": [deepcopy(ARGUMENT)]}
                del payload["arguments"][0]["sources"][0][field]
                self.assertEqual(self.run_export(payload)[0], 2)
                self.assert_no_output()

    def test_rejects_unknown_sources_title_only_records_and_non_http_urls(self):
        invalid_sources = [
            {"title": "只有一个标题"},
            {**ARGUMENT["sources"][0], "author": "未知来源"},
        ]
        invalid_sources += [
            {**ARGUMENT["sources"][0], "url": url}
            for url in ("ftp://example.org/source", "javascript:alert(1)", "https://", "https://example.org/a b", "https://example.org:bad/a")
        ]
        for source in invalid_sources:
            with self.subTest(source=source):
                payload = {"arguments": [deepcopy(ARGUMENT)]}
                payload["arguments"][0]["sources"] = [source]
                self.assertEqual(self.run_export(payload)[0], 2)
                self.assert_no_output()

    def test_rejects_duplicate_ids_and_invalid_categories(self):
        payload = self.payload()
        payload["arguments"][1]["id"] = "A1"
        self.assertEqual(self.run_export(payload)[0], 2)
        self.assert_no_output()
        for categories in ([], ["攻防"], ["进攻", "进攻"], "进攻"):
            with self.subTest(categories=categories):
                payload = {"arguments": [deepcopy(ARGUMENT)]}
                payload["arguments"][0]["categories"] = categories
                self.assertEqual(self.run_export(payload)[0], 2)
                self.assert_no_output()

    def test_refuses_same_name_without_overwriting_and_force_is_explicit(self):
        self.assertEqual(self.run_export(self.payload())[0], 0)
        output = self.output_path()
        original_bytes = output.read_bytes()
        changed = self.payload()
        changed["arguments"][0]["claim"] = "显式覆盖后的测试论据"
        with patch.object(exporter, "build_argument_document") as builder:
            code, _, _ = self.run_export(changed)
        self.assertEqual(code, 3)
        builder.assert_not_called()
        self.assertEqual(output.read_bytes(), original_bytes)
        code, _, stderr = self.run_export(changed, "--force")
        self.assertEqual(code, 0, stderr)
        doc = Document(output)
        self.assertIn("显式覆盖后的测试论据", doc.tables[0].rows[1].cells[1].text)

    def test_custom_directory_is_flat_and_filename_uses_full_topic_cleaning(self):
        self.assertEqual(self.run_export(self.payload())[0], 0)
        self.assertEqual(list(self.products.iterdir()), [self.output_path()])
        topic = ' 教育/科技:比较? '
        self.assertEqual(exporter.argument_table_filename(topic, "正方"), "教育-科技-比较--正方-论据资料表.docx")
        self.assertEqual(exporter.argument_table_filename("CON.txt", "正方"), "_CON.txt-正方-论据资料表.docx")

    def test_default_location_reuses_desktop_products_without_subdirectory(self):
        desktop_products = self.root / "redirected desktop" / "products"
        with patch.object(exporter, "default_products_dir", return_value=desktop_products) as desktop:
            code, stdout, stderr = self.run_export(self.payload(), custom_directory=False)
        self.assertEqual(code, 0, stderr)
        desktop.assert_called_once_with()
        self.assertTrue(self.output_path(desktop_products).is_file())
        self.assertEqual(list(desktop_products.iterdir()), [self.output_path(desktop_products)])
        self.assertIn("桌面上的products文件夹", stdout)

    def test_empty_arguments_and_invalid_json_create_no_product_directory(self):
        code, stdout, _ = self.run_export({"arguments": []})
        self.assertEqual(code, 2)
        self.assertNotIn("已保存", stdout)
        self.assert_no_output()
        self.input_file.write_text("not JSON", encoding="utf-8")
        self.assertEqual(self.run_export()[0], 2)
        self.assert_no_output()


if __name__ == "__main__":
    unittest.main()
