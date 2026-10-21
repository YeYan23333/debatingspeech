"""Offer an optional evidence table only after verified DOCX delivery."""
from __future__ import annotations

import unittest
from unittest.mock import patch

from docx import Document

try:
    from . import test_desktop_output as desktop_fixtures
except ImportError:
    import test_desktop_output as desktop_fixtures


exporter = desktop_fixtures.exporter
generator = desktop_fixtures.generator
QUESTION = "是否将本次检索中已核验的论据，按“进攻、防守、证明”分类整理成资料表？"


class ArgumentTableOfferTests(unittest.TestCase):
    # Reuse CLI fixtures without inheriting and rerunning their test methods.
    setUp = desktop_fixtures.DesktopOutputTests.setUp
    output_path = desktop_fixtures.DesktopOutputTests.output_path
    run_export = desktop_fixtures.DesktopOutputTests.run_export
    run_generator = desktop_fixtures.DesktopOutputTests.run_generator

    def assert_offer_follows_success(self, stdout, side):
        lines = stdout.splitlines()
        closing = exporter.completion_message(self.products, default_location=True)
        self.assertEqual(lines[-2:], [closing, QUESTION])
        self.assertEqual(lines.count(QUESTION), 1)
        document = Document(self.output_path(side))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("我方认为应该重视学习方式的变化", text)
        self.assertNotIn(QUESTION, text)

    def test_export_success_offers_table_on_separate_line_without_waiting(self):
        with patch("builtins.input", side_effect=AssertionError("CLI must not wait for a reply")):
            code, stdout, stderr = self.run_export()
        self.assertEqual(code, 0, stderr)
        self.assert_offer_follows_success(stdout, "正方")

    def test_model_success_offers_table_without_waiting_or_extra_model_request(self):
        with patch("builtins.input", side_effect=AssertionError("CLI must not wait for a reply")):
            code, stdout, stderr, _, router = self.run_generator()
        self.assertEqual(code, 0, stderr)
        router.generate.assert_called_once()
        self.assert_offer_follows_success(stdout, "反方")

    def test_dry_run_does_not_offer_table_or_call_model(self):
        code, stdout, stderr, router_class, router = self.run_generator("--dry-run")
        self.assertEqual(code, 0, stderr)
        self.assertNotIn(QUESTION, stdout + stderr)
        self.assertFalse(self.products.exists())
        router_class.assert_not_called()
        router.generate.assert_not_called()

    def test_empty_export_does_not_offer_table(self):
        self.input_file.write_text(" \n", encoding="utf-8")
        code, stdout, stderr = self.run_export()
        self.assertNotEqual(code, 0)
        self.assertNotIn(QUESTION, stdout + stderr)

    def test_model_failure_does_not_offer_table(self):
        code, stdout, stderr, _, router = self.run_generator(
            model_failure=generator.ProviderError("模拟离线错误")
        )
        self.assertNotEqual(code, 0)
        self.assertNotIn(QUESTION, stdout + stderr)
        router.generate.assert_called_once()

    def test_export_refusal_to_overwrite_does_not_offer_table(self):
        self.assertEqual(self.run_export()[0], 0)
        before = self.output_path().read_bytes()
        code, stdout, stderr = self.run_export()
        self.assertEqual(code, 3)
        self.assertNotIn(QUESTION, stdout + stderr)
        self.assertEqual(self.output_path().read_bytes(), before)

    def test_model_refusal_to_overwrite_does_not_offer_table(self):
        self.assertEqual(self.run_generator()[0], 0)
        before = self.output_path("反方").read_bytes()
        code, stdout, stderr, router_class, router = self.run_generator()
        self.assertEqual(code, 3)
        self.assertNotIn(QUESTION, stdout + stderr)
        self.assertEqual(self.output_path("反方").read_bytes(), before)
        router_class.assert_not_called()
        router.generate.assert_not_called()

    def test_export_reopen_failure_does_not_offer_table(self):
        # The first Document() builds the file; the second is its reopen check.
        with patch.object(exporter, "Document", side_effect=[Document(), OSError("模拟重开失败")]):
            code, stdout, stderr = self.run_export()
        self.assertNotEqual(code, 0)
        self.assertTrue(self.output_path().exists())
        self.assertNotIn(QUESTION, stdout + stderr)

    def test_model_reopen_failure_does_not_offer_table(self):
        # The exporter already holds its real Document constructor; this patches
        # the generator's later local import used to reopen the saved document.
        with patch("docx.Document", side_effect=OSError("模拟重开失败")):
            code, stdout, stderr, _, _ = self.run_generator()
        self.assertNotEqual(code, 0)
        self.assertTrue(self.output_path("反方").exists())
        self.assertNotIn(QUESTION, stdout + stderr)


if __name__ == "__main__":
    unittest.main()
