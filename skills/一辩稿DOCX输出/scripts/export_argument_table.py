#!/usr/bin/env python3
"""Export an explicitly requested argument collection as real Word tables.

The input must contain arguments with verified=true and public-source verification
notes. This script validates the JSON structure and declared status only: it does
not use the network, verify facts, or turn unknown/title-only sources into verified
evidence. verification_note describes checks of public evidence and any secondary
citation; it must not contain private reasoning or internal deliberation.
"""
from __future__ import annotations

import argparse
from io import BytesIO
import json
import os
from pathlib import Path
import re
import sys
import tempfile
from urllib.parse import urlsplit

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.text.run import Run
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from export_to_docx import default_products_dir, set_run_font, set_style_font, speech_filename


CATEGORIES = ("进攻", "防守", "证明")
NO_MATERIAL_MESSAGE = "本次未找到可核验的相关论据"
HEADERS = ("编号", "论据与材料", "用途", "适用条件", "核验说明与来源")
COLUMN_WIDTHS_CM = (1.5, 6.5, 4.8, 4.3, 9.6)
TEXT_FIELDS = ("claim", "material", "usage", "limitation", "verification_note")
SOURCE_FIELDS = ("author", "title", "date", "url")
UNKNOWN_SOURCE_VALUES = frozenset({
    "未知", "未知来源", "来源未知", "不详", "暂无", "待核验", "未核验", "未查",
    "无出处", "无来源", "未提供", "待补充", "仅有标题", "unknown", "n/a", "none", "null",
})


def _text(value, field: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field} 必须是非空字符串。")
    text = value.strip()
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", text):
        raise ValueError(f"{field} 含有 Word 不支持的控制字符。")
    return text


def _source_url(value, field: str) -> str:
    url = _text(value, field)
    try:
        parts = urlsplit(url)
        valid = parts.scheme.lower() in {"http", "https"} and bool(parts.netloc) and bool(parts.hostname)
        # Accessing port also rejects malformed bracket/port syntax.
        parts.port
    except ValueError:
        valid = False
    if not valid or re.search(r"\s", url):
        raise ValueError(f"{field} 必须是完整的 http/https 来源链接。")
    return url


def validate_payload(payload) -> list[dict]:
    """Validate declarations without treating that validation as fact-checking."""
    if not isinstance(payload, dict) or not isinstance(payload.get("arguments"), list):
        raise ValueError("JSON 顶层必须是包含 arguments 数组的对象。")
    if not payload["arguments"]:
        raise ValueError("arguments 为空；本次没有可导出的已核验相关论据。")
    normalized = []
    seen_ids = set()
    for index, item in enumerate(payload["arguments"], 1):
        prefix = f"arguments[{index}]"
        if not isinstance(item, dict):
            raise ValueError(f"{prefix} 必须是对象。")
        identifier = item.get("id")
        if isinstance(identifier, int) and not isinstance(identifier, bool):
            identifier = str(identifier)
        identifier = _text(identifier, f"{prefix}.id")
        if identifier in seen_ids:
            raise ValueError(f"论据 id 重复：{identifier}")
        seen_ids.add(identifier)
        if item.get("verified") is not True:
            raise ValueError(f"{prefix}.verified 必须为布尔 true；待核验条目不能导出。")
        categories = item.get("categories")
        if (not isinstance(categories, list) or not categories
                or any(not isinstance(category, str) or category not in CATEGORIES for category in categories)
                or len(set(categories)) != len(categories)):
            raise ValueError(f"{prefix}.categories 必须包含进攻、防守、证明中的一个或多个不同类别。")
        result = {"id": identifier, "verified": True, "categories": list(categories)}
        for field in TEXT_FIELDS:
            result[field] = _text(item.get(field), f"{prefix}.{field}")
        sources = item.get("sources")
        if not isinstance(sources, list) or not sources:
            raise ValueError(f"{prefix}.sources 必须是非空来源数组，不能只有标题或未知出处。")
        result["sources"] = []
        for source_index, source in enumerate(sources, 1):
            source_prefix = f"{prefix}.sources[{source_index}]"
            if not isinstance(source, dict):
                raise ValueError(f"{source_prefix} 必须是对象。")
            checked = {}
            for field in SOURCE_FIELDS:
                text = _text(source.get(field), f"{source_prefix}.{field}")
                if text.lower() in UNKNOWN_SOURCE_VALUES:
                    raise ValueError(f"{source_prefix}.{field} 不能使用未知或待核验占位内容。")
                checked[field] = text
            checked["url"] = _source_url(checked["url"], f"{source_prefix}.url")
            result["sources"].append(checked)
        normalized.append(result)
    return normalized


def argument_table_filename(topic: str, side: str) -> str:
    """Keep the full-topic filename sanitation used by speech exports."""
    filename = speech_filename(topic, side)
    return filename[:-len("一辩稿.docx")] + "论据资料表.docx"


def _cell_paragraph(cell, text: str, *, bold: bool = False, first: bool = False):
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.1
    paragraph.paragraph_format.keep_with_next = False
    # Permit long Latin URLs to wrap without inserting characters into the URL.
    word_wrap = OxmlElement("w:wordWrap")
    word_wrap.set(qn("w:val"), "0")
    paragraph._p.get_or_add_pPr().append(word_wrap)
    run = paragraph.add_run(text)
    set_run_font(run, east_asia="宋体", latin="Arial")
    run.font.size = Pt(10)
    run.bold = bold
    return paragraph


def _cell_hyperlink(cell, url: str):
    """Keep the full URL readable while making it a real external Word link."""
    paragraph = _cell_paragraph(cell, "")
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    hyperlink.set(qn("w:history"), "1")
    run_element = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = url
    run_element.append(text)
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)
    run = Run(run_element, paragraph)
    set_run_font(run, east_asia="宋体", latin="Arial")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string("0563C1")
    run.font.underline = True
    return paragraph


def _format_table(table):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, COLUMN_WIDTHS_CM):
        column.width = Cm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, COLUMN_WIDTHS_CM):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{side}")
                margin.set(qn("w:w"), "90")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
    for cell, heading in zip(table.rows[0].cells, HEADERS):
        _cell_paragraph(cell, heading, bold=True, first=True)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E8EDF2")
        cell._tc.get_or_add_tcPr().append(shading)


def build_argument_document(arguments: list[dict], topic: str, side: str):
    """Build native Word tables; validation must precede this function."""
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.top_margin, section.bottom_margin = Cm(1.5), Cm(1.5)
    section.left_margin, section.right_margin = Cm(1.5), Cm(1.5)
    set_style_font(doc.styles["Normal"], "宋体", "Arial")
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    for style_name in ("Title", "Heading 1"):
        style = doc.styles[style_name]
        set_style_font(style, "黑体", "Arial")
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(13)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    doc.core_properties.title = f"{topic} {side} 论据资料表"
    doc.core_properties.subject = "华语辩论已核验论据资料表"
    doc.core_properties.keywords = "辩论,论据,进攻,防守,证明"
    title = doc.add_paragraph("论据资料表", style="Title")
    title.paragraph_format.space_after = Pt(6)
    doc.add_paragraph(f"辩题：{topic}    立场：{side}")
    doc.add_paragraph("按进攻、防守、证明分类。同一编号出现在不同类别时，指向同一条论据。")
    for category in CATEGORIES:
        doc.add_paragraph(category, style="Heading 1")
        selected = [argument for argument in arguments if category in argument["categories"]]
        if not selected:
            doc.add_paragraph(NO_MATERIAL_MESSAGE)
            continue
        table = doc.add_table(rows=1 + len(selected), cols=len(HEADERS))
        _format_table(table)
        for row, argument in zip(table.rows[1:], selected):
            _cell_paragraph(row.cells[0], argument["id"], first=True)
            _cell_paragraph(row.cells[1], argument["claim"], bold=True, first=True)
            _cell_paragraph(row.cells[1], "材料：" + argument["material"])
            _cell_paragraph(row.cells[2], argument["usage"], first=True)
            _cell_paragraph(row.cells[3], argument["limitation"], first=True)
            _cell_paragraph(row.cells[4], argument["verification_note"], first=True)
            for number, source in enumerate(argument["sources"], 1):
                _cell_paragraph(
                    row.cells[4],
                    f"[{number}] {source['author']}　{source['title']}（{source['date']}）",
                )
                _cell_hyperlink(row.cells[4], source["url"])
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return doc


def check_saved_table(output_path: Path):
    doc = Document(output_path)
    if not any(len(table.rows) > 1 and any(cell.text.strip() for cell in table.rows[1].cells)
               for table in doc.tables):
        raise RuntimeError("生成的 DOCX 未包含非空论据表格。")


def save_argument_document(doc, output_path: Path, *, force: bool):
    """Validate a saved temporary document and refuse accidental replacement."""
    buffer = BytesIO()
    doc.save(buffer)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(suffix=".docx", prefix=".argument-table-", dir=output_path.parent)
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "wb") as temporary_file:
            temporary_file.write(buffer.getvalue())
        check_saved_table(temporary_path)
        if force:
            os.replace(temporary_path, output_path)
        else:
            # Exclusive creation also protects a file created after the CLI check.
            created = False
            try:
                with output_path.open("xb") as output_file:
                    created = True
                    output_file.write(buffer.getvalue())
                check_saved_table(output_path)
            except Exception:
                if created:
                    output_path.unlink(missing_ok=True)
                raise
        check_saved_table(output_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def table_completion_message(products_dir: Path, *, default_location: bool) -> str:
    location = "桌面上的products" if default_location else str(products_dir.resolve())
    return f"论据资料表已保存，请在{location}文件夹中检查。"


def main() -> int:
    parser = argparse.ArgumentParser(
        description="将用户选择生成的已核验论据 JSON 导出为 Word 资料表。",
        epilog=("本脚本仅校验字段结构与 verified=true 声明，不联网、不自动核实事实。"
                "不得将仅有标题、未知出处或待核验条目标为已核验。verification_note 只记公开证据"
                "的核对结果和转引情况，不记录内部思考链。"),
    )
    parser.add_argument("--topic", required=True, help="用户询问的完整辩题")
    parser.add_argument("--side", required=True, choices=["正方", "反方"], help="立场")
    parser.add_argument("--input", required=True, help="UTF-8 JSON 文件，顶层含 arguments 数组")
    parser.add_argument("--products-dir", help="成品目录，默认系统桌面的 products；不建立辩题子目录")
    parser.add_argument("--force", action="store_true", help="明确允许覆盖同名资料表")
    args = parser.parse_args()
    try:
        payload = json.loads(Path(args.input).read_text(encoding="utf-8-sig"))
        arguments = validate_payload(payload)
        products_dir = Path(args.products_dir).expanduser() if args.products_dir else default_products_dir()
        output_path = products_dir / argument_table_filename(args.topic, args.side)
        if output_path.exists() and not args.force:
            print(f"目标资料表已存在，为避免覆盖已停止：{output_path}。确认覆盖请增加 --force。", file=sys.stderr)
            return 3
        document = build_argument_document(arguments, args.topic, args.side)
        save_argument_document(document, output_path, force=args.force)
        print(output_path.resolve())
        print(table_completion_message(products_dir, default_location=args.products_dir is None))
        return 0
    except FileExistsError:
        print("目标资料表已存在，未覆盖；确认覆盖请增加 --force。", file=sys.stderr)
        return 3
    except (OSError, ValueError, RuntimeError) as exc:
        print(f"资料表导出失败：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
